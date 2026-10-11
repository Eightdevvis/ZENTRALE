# core/tabelle_datei.py
#
# Excel-Tabellen (.xlsx) lesen, anlegen und ändern; CSV lesen und schreiben.
# Für den Skill excel (memory/ki/tabellen.md): Sasha will Kosten und Stunden
# in einer Tabelle führen, die wächst und selbst rechnet — ein PDF kann das
# nicht (Gespräch 20261011-081303).
#
# 2026-10-11. Bibliothek: openpyxl (MIT, reines Python, dazu et-xmlfile) —
# läuft auch auf dem Pi ohne Bauen. Sie RECHNET NICHT: beim Anlegen stehen
# Formeln als Formeln in der Datei (Excel/LibreOffice rechnen beim Öffnen,
# fullCalcOnLoad ist gesetzt); beim Lesen einer von Excel gespeicherten Datei
# kommt der dort gespeicherte Wert mit, sonst steht „kein gespeicherter Wert".
#
# ── Warum in einem eigenen Prozess ─────────────────────────────────────
# Wie bei PDFs (core/pdf_datei.py): eine xlsx ist eine Zip voller XML und
# kommt oft von draußen. Zip-Bomben prüft der Elternprozess vorher (Zahl und
# entpackte Größe der Einträge); openpyxl läuft im Kindprozess dieser Datei
# (`python -I tabelle_datei.py`) mit Zeit-, CPU- und Speichergrenze. Der
# Kind-Teil (unten) importiert nur openpyxl und die Standardbibliothek.
#
# CSV ist Text und läuft hier im Prozess (csv-Modul), mit Grenzen.
#
# Dienst (Schicht 2, memory/system/bauplan_kern.md).

import csv
import io
import json
import os
import resource
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, time

MAX_BYTES = 30 * 1024 * 1024
MAX_ENTPACKT = 200 * 1024 * 1024     # gegen Zip-Bomben
MAX_EINTRAEGE = 5000
ZEIT_S = 60
SPEICHER_BYTES = 1536 * 1024 * 1024
MAX_BLAETTER = 50
MAX_ZELLEN = 100_000                 # was ein Aufruf anlegen/ändern darf
MAX_ZELLE_ZEICHEN = 32_767           # Excel-Grenze je Zelle
LESEN_ZELLEN = 2_000                 # so viele gefüllte Zellen je Lesen

# Was openpyxl beim Ändern nicht mitnimmt — dann gibt es keine geänderte
# Fassung (ganz oder gar nicht: eine Kopie ohne Sashas Diagramm wäre „anders
# als verlangt").
VERLUST = {"xl/vbaProject.bin": "Makros", "xl/pivotTables/": "Pivot-Tabellen",
           "xl/drawings/": "Diagramme oder Bilder", "xl/externalLinks/": "Verknüpfungen "
           "zu anderen Dateien", "xl/slicers/": "Datenschnitte"}

FEHLER_TEXT = {
    "kaputt": "Die Datei ist keine lesbare Excel-Tabelle (.xlsx) — beschädigt oder etwas anderes.",
    "gesperrt": "Die Tabelle ist mit einem Passwort geschützt — ohne komme ich nicht hinein.",
    "bibliothek": ("Die Excel-Bibliothek (openpyxl) fehlt auf diesem Rechner — "
                   "venv/bin/pip install -r requirements.txt"),
    "zeit": "Die Tabelle hat zu lange gebraucht (abgebrochen) — vermutlich kaputt oder riesig.",
}


class Fehler(ValueError):
    """Etwas ging mit der Tabelle nicht. art: kaputt | gesperrt | bibliothek |
    zeit | zu_gross | blatt | bereich | zelle | verlust (für den Fehlercode)."""

    def __init__(self, text: str, art: str = "kaputt"):
        super().__init__(text)
        self.art = art


# ── Adressen ────────────────────────────────────────────────────────────

def spalte(n: int) -> str:
    """1 → A, 27 → AA."""
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def adresse(zeile: int, sp: int) -> str:
    return f"{spalte(sp)}{zeile}"


def zelle_zerlegen(text: str) -> tuple:
    """„B12" → (12, 2). Wirft Fehler (art zelle)."""
    t = str(text or "").strip().upper().replace("$", "")
    i = 0
    while i < len(t) and t[i].isalpha():
        i += 1
    buchst, ziffern = t[:i], t[i:]
    if not (1 <= len(buchst) <= 3) or not ziffern.isdigit() or ziffern.startswith("0"):
        raise Fehler(f"„{text}“ ist keine Zelle (so: B12).", "zelle")
    sp = 0
    for b in buchst:
        sp = sp * 26 + (ord(b) - 64)
    z = int(ziffern)
    if sp > 16384 or z > 1_048_576:
        raise Fehler(f"„{text}“ liegt außerhalb einer Excel-Tabelle.", "zelle")
    return z, sp


def bereich_zerlegen(text) -> tuple | None:
    """„A1:D20" → (1, 1, 20, 4); „B5" → eine Zelle; leer → None."""
    t = str(text or "").strip()
    if not t:
        return None
    teile = t.split(":")
    if len(teile) > 2:
        raise Fehler(f"„{text}“ ist kein Bereich (so: A1:D20).", "bereich")
    try:
        z1, s1 = zelle_zerlegen(teile[0])
        z2, s2 = zelle_zerlegen(teile[-1])
    except Fehler:
        raise Fehler(f"„{text}“ ist kein Bereich (so: A1:D20).", "bereich")
    return min(z1, z2), min(s1, s2), max(z1, z2), max(s1, s2)


# ── Werte ───────────────────────────────────────────────────────────────

def wert_text(w) -> str:
    """Wie ein Wert für die KI aussieht: Text in Anführung, Zahl blank."""
    if w is None:
        return "(leer)"
    if isinstance(w, bool):
        return "WAHR" if w else "FALSCH"
    if isinstance(w, float) and w.is_integer() and abs(w) < 1e15:
        return str(int(w))
    if isinstance(w, (int, float)):
        return str(w)
    return json.dumps(str(w), ensure_ascii=False)


# ── xlsx erkennen, Zip prüfen ───────────────────────────────────────────

def ist_xlsx(daten: bytes) -> bool:
    if not daten or daten[:4] != b"PK\x03\x04":
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(daten)) as z:
            return "xl/workbook.xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


def _zip_pruefen(daten: bytes) -> list:
    """Namen der Einträge — oder Fehler (zu groß, Bombe, kaputt)."""
    if not daten:
        raise Fehler("Die Datei ist leer.")
    if len(daten) > MAX_BYTES:
        raise Fehler(f"Zu groß — höchstens {MAX_BYTES // 1024 // 1024} MB.", "zu_gross")
    if daten[:4] == b"\xd0\xcf\x11\xe0":
        raise Fehler("Eine alte Excel-Datei (.xls) oder mit Passwort — bitte als .xlsx "
                     "ohne Passwort speichern.", "gesperrt")
    try:
        with zipfile.ZipFile(io.BytesIO(daten)) as z:
            infos = z.infolist()
    except zipfile.BadZipFile:
        raise Fehler(FEHLER_TEXT["kaputt"])
    if len(infos) > MAX_EINTRAEGE or sum(i.file_size for i in infos) > MAX_ENTPACKT:
        raise Fehler("Entpackt viel zu groß (Zip-Bombe?) — nicht geöffnet.", "zu_gross")
    namen = [i.filename for i in infos]
    if "xl/workbook.xml" not in namen:
        raise Fehler(FEHLER_TEXT["kaputt"])
    return namen


def verlust(daten: bytes) -> list:
    """Was beim Ändern verloren ginge (Wörter für Sasha)."""
    namen = _zip_pruefen(daten)
    return sorted({wort for teil, wort in VERLUST.items()
                   if any(n == teil or n.startswith(teil) for n in namen)})


# ── Kindprozess starten ─────────────────────────────────────────────────

def _grenzen():
    for art, wert in ((resource.RLIMIT_AS, SPEICHER_BYTES),
                      (resource.RLIMIT_CPU, int(ZEIT_S) + 5),
                      (resource.RLIMIT_FSIZE, 4 * MAX_BYTES)):
        try:
            resource.setrlimit(art, (wert, wert))
        except (ValueError, OSError):
            pass


def _im_kind(auftrag: dict, eingabe: bytes | None = None) -> tuple:
    """auftrag (+ xlsx-Bytes) → (antwort, ausgabe_bytes|None). Wirft Fehler."""
    if eingabe is not None:
        _zip_pruefen(eingabe)
    with tempfile.TemporaryDirectory(prefix="zentrale-xlsx-") as tmp:
        auftrag = dict(auftrag, ziel=os.path.join(tmp, "aus.xlsx"))
        if eingabe is not None:
            auftrag["datei"] = os.path.join(tmp, "ein.xlsx")
            with open(auftrag["datei"], "wb") as f:
                f.write(eingabe)
        try:
            proc = subprocess.run(
                [sys.executable, "-I", "-B", os.path.abspath(__file__)],
                input=json.dumps(auftrag, ensure_ascii=False), capture_output=True, text=True,
                timeout=ZEIT_S, preexec_fn=_grenzen, cwd=tmp,
                env={"PATH": "/usr/bin:/bin", "LANG": "C.UTF-8"})
        except subprocess.TimeoutExpired:
            raise Fehler(FEHLER_TEXT["zeit"], "zeit")
        except (OSError, subprocess.SubprocessError):
            raise Fehler("Der Excel-Leser ließ sich nicht starten.")
        try:
            antwort = json.loads(proc.stdout or "{}")
        except ValueError:
            antwort = {}
        if not antwort:
            raise Fehler(FEHLER_TEXT["zeit"] if proc.returncode < 0 else FEHLER_TEXT["kaputt"],
                         "zeit" if proc.returncode < 0 else "kaputt")
        if antwort.get("fehler"):
            art = antwort["fehler"]
            # blatt/bereich bringen einen eigenen Satz mit (aus diesem Modul);
            # alles andere nur als Art, nie Rohes aus der Bibliothek.
            text = antwort.get("meldung") if art in ("blatt", "bereich", "zelle") \
                else FEHLER_TEXT.get(art, FEHLER_TEXT["kaputt"])
            raise Fehler(text or FEHLER_TEXT["kaputt"], art)
        aus = None
        if os.path.isfile(auftrag["ziel"]):
            with open(auftrag["ziel"], "rb") as f:
                aus = f.read()
        return antwort, aus


# ── Öffentlich: xlsx ────────────────────────────────────────────────────

def lesen(daten: bytes, blatt=None, bereich=None) -> dict:
    """→ {blaetter: [{name, bereich, zeilen, spalten}], blatt, bereich,
    zellen: [[adresse, wert, formel|None, berechnet|None]], weiter: Bereich|None}.
    Ohne Bereich ab A1; je Aufruf höchstens LESEN_ZELLEN gefüllte Zellen."""
    b = bereich_zerlegen(bereich)
    return _im_kind({"aufgabe": "lesen", "blatt": blatt, "bereich": b}, daten)[0]


def erzeugen(blaetter: list) -> bytes:
    """[{name, zeilen: [[wert]], spaltenbreiten?: [zahl]}] → xlsx-Bytes.
    Werte schon geprüft (ki_tabelle): Zahl, Text, WAHR/FALSCH, None, Formel
    als Text mit = (englisch)."""
    return _im_kind({"aufgabe": "anlegen", "blaetter": blaetter})[1]


def aendern(daten: bytes, aenderungen: list, anhaengen: list) -> tuple:
    """Eine geänderte Fassung. aenderungen: [{blatt, zelle, wert}], anhaengen:
    [{blatt, zeilen}] (blatt None = das erste). → (bytes, {angehaengt: [{blatt,
    von, bis}], blaetter}). Wirft Fehler (auch art verlust)."""
    weg = verlust(daten)
    if weg:
        raise Fehler(f"Beim Ändern gingen verloren: {', '.join(weg)} — diese Tabelle "
                     f"ändert Sasha besser selbst.", "verlust")
    antwort, aus = _im_kind({"aufgabe": "aendern", "aenderungen": aenderungen,
                             "anhaengen": anhaengen}, daten)
    return aus, antwort


# ── CSV ─────────────────────────────────────────────────────────────────

CSV_MAX_ZEICHEN = 2_000_000


def csv_lesen(text: str) -> tuple:
    """→ (zeilen [[str]], trenner). Komma, Semikolon oder Tab: was in den
    ersten Zeilen am gleichmäßigsten vorkommt (Sniffer rät sonst gern „,"
    bei deutschen Zahlen mit Komma)."""
    text = str(text or "")
    if len(text) > CSV_MAX_ZEICHEN:
        raise Fehler("Die CSV ist zu lang.", "zu_gross")
    text = text.lstrip("﻿")
    probe = [z for z in text.splitlines()[:20] if z.strip()]
    bester, wert = ",", -1.0
    for t in (";", ",", "\t"):
        zahlen = [z.count(t) for z in probe]
        if not zahlen or not zahlen[0]:
            continue
        gleich = sum(1 for n in zahlen if n == zahlen[0]) / len(zahlen)
        if gleich + zahlen[0] / 1000 > wert:
            bester, wert = t, gleich + zahlen[0] / 1000
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=bester)], bester


def csv_schreiben(zeilen: list, trenner: str = ";") -> str:
    """Zeilen → CSV-Text. Semikolon, weil Sashas Excel deutsch ist (dort
    trennt „;"; Komma ist das Dezimalzeichen)."""
    aus = io.StringIO()
    w = csv.writer(aus, delimiter=trenner, lineterminator="\n")
    for z in zeilen:
        w.writerow(["" if v is None else _csv_wert(v, trenner) for v in z])
    return aus.getvalue()


def _csv_wert(v, trenner):
    if isinstance(v, bool):
        return "WAHR" if v else "FALSCH"
    if isinstance(v, float) and trenner == ";":
        return (str(int(v)) if v.is_integer() else str(v)).replace(".", ",")
    return str(v)


def csv_zellen(zeilen: list, bereich=None) -> dict:
    """CSV-Zeilen in dieselbe Form wie lesen() (ein Blatt „CSV")."""
    b = bereich_zerlegen(bereich)
    z1, s1, z2, s2 = b or (1, 1, 10 ** 9, 10 ** 9)
    zellen, weiter = [], None
    breite = max((len(r) for r in zeilen), default=0)
    for zi, reihe in enumerate(zeilen, start=1):
        if zi < z1 or zi > z2:
            continue
        gefuellt = [(si, v) for si, v in enumerate(reihe, start=1)
                    if s1 <= si <= s2 and str(v).strip() != ""]
        if len(zellen) + len(gefuellt) > LESEN_ZELLEN and zellen:
            weiter = f"{spalte(s1)}{zi}:{spalte(min(s2, max(breite, 1)))}{min(z2, len(zeilen))}"
            break
        zellen += [[adresse(zi, si), v, None, None] for si, v in gefuellt]
    dim = f"A1:{spalte(max(breite, 1))}{max(len(zeilen), 1)}"
    return {"blaetter": [{"name": "CSV", "bereich": dim, "zeilen": len(zeilen),
                          "spalten": breite}], "blatt": "CSV",
            "bereich": bereich or dim, "zellen": zellen, "weiter": weiter}


# ── Als Text (für die KI und die Vorschau) ──────────────────────────────

def als_text(r: dict) -> str:
    """Zeilenweise: „2: A2 "Grundbetrag" | B2 380 | B14 =SUM(B2:B13) → 440"."""
    zeilen, aktuell, teile = [], None, []
    for adr, wert, formel, berechnet in r.get("zellen") or []:
        nr = int("".join(c for c in adr if c.isdigit()))
        if nr != aktuell and teile:
            zeilen.append(f"{aktuell}: " + " | ".join(teile))
            teile = []
        aktuell = nr
        if formel:
            f = str(formel).replace("_xlfn.", "")
            dahinter = (f" → {wert_text(berechnet)}" if berechnet is not None
                        else " → (kein gespeicherter Wert)")
            teile.append(f"{adr} {f}{dahinter}")
        else:
            teile.append(f"{adr} {wert_text(wert)}")
    if teile:
        zeilen.append(f"{aktuell}: " + " | ".join(teile))
    return "\n".join(zeilen)


def kopf_text(r: dict) -> str:
    bl = ", ".join(f"{b['name']} ({b['bereich']})" for b in r.get("blaetter") or [])
    return f"Blätter: {bl}." if bl else "Keine Blätter."


# ── Im Kindprozess ──────────────────────────────────────────────────────
# Nur Standardbibliothek + openpyxl. Antwort als JSON auf stdout; Fehler als
# {"fehler": art} — nie eine Python-Meldung mit Pfaden oder Inhalten.

def _kind_wert(v):
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    t = getattr(v, "text", None)            # ArrayFormula & Co.
    return str(t if t is not None else v)


def _kind_blatt(wb, name):
    if name in (None, ""):
        return wb.worksheets[0]
    for ws in wb.worksheets:
        if ws.title.lower() == str(name).strip().lower():
            return ws
    raise Fehler(f"Ein Blatt „{name}“ gibt es nicht (da sind: "
                 + ", ".join(ws.title for ws in wb.worksheets) + ").", "blatt")


def _kind_lesen(auftrag):
    from openpyxl import load_workbook
    wb_f = load_workbook(auftrag["datei"], read_only=True, data_only=False)
    wb_w = load_workbook(auftrag["datei"], read_only=True, data_only=True)
    blaetter = []
    for ws in wb_f.worksheets:
        try:
            dim = ws.calculate_dimension()
        except Exception:
            dim = "A1:A1"
        blaetter.append({"name": ws.title, "bereich": dim, "zeilen": ws.max_row or 0,
                         "spalten": ws.max_column or 0})
    ws_f = _kind_blatt(wb_f, auftrag.get("blatt"))
    ws_w = wb_w[ws_f.title]
    b = auftrag.get("bereich")
    z1, s1, z2, s2 = b or (1, 1, max(ws_f.max_row or 1, 1), max(ws_f.max_column or 1, 1))
    werte = {}
    for zi, reihe in enumerate(ws_w.iter_rows(min_row=z1, max_row=z2, min_col=s1, max_col=s2,
                                              values_only=True), start=z1):
        for si, v in enumerate(reihe, start=s1):
            if v is not None:
                werte[(zi, si)] = _kind_wert(v)
    zellen, weiter = [], None
    for zi, reihe in enumerate(ws_f.iter_rows(min_row=z1, max_row=z2, min_col=s1, max_col=s2,
                                              values_only=True), start=z1):
        neu = []
        for si, v in enumerate(reihe, start=s1):
            if v is None or (isinstance(v, str) and v == ""):
                continue
            v = _kind_wert(v)
            if isinstance(v, str) and v.startswith("="):
                neu.append([adresse(zi, si), None, v, werte.get((zi, si))])
            else:
                neu.append([adresse(zi, si), v, None, None])
        if zellen and len(zellen) + len(neu) > LESEN_ZELLEN:
            weiter = f"{spalte(s1)}{zi}:{spalte(s2)}{z2}"
            break
        zellen += neu
    return {"blaetter": blaetter, "blatt": ws_f.title,
            "bereich": f"{spalte(s1)}{z1}:{spalte(s2)}{z2}", "zellen": zellen, "weiter": weiter}


def _kind_anlegen(auftrag):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for bl in auftrag["blaetter"]:
        ws = wb.create_sheet(bl["name"])
        for reihe in bl["zeilen"]:
            ws.append(reihe)
        for i, breite in enumerate(bl.get("spaltenbreiten") or [], start=1):
            if breite:
                ws.column_dimensions[get_column_letter(i)].width = float(breite)
    wb.calculation.fullCalcOnLoad = True
    wb.save(auftrag["ziel"])
    return {"blaetter": [b["name"] for b in auftrag["blaetter"]]}


def _kind_aendern(auftrag):
    from openpyxl import load_workbook
    wb = load_workbook(auftrag["datei"])
    for a in auftrag.get("aenderungen") or []:
        ws = _kind_blatt(wb, a.get("blatt"))
        ws[a["zelle"]] = a.get("wert")
    angehaengt = []
    for a in auftrag.get("anhaengen") or []:
        ws = _kind_blatt(wb, a.get("blatt"))
        # Die letzte Zeile mit Inhalt — nicht max_row (formatierte Leerzeilen
        # zählen dort mit, angehängt würde sonst weit unten).
        letzte = 0
        for reihe in ws.iter_rows():
            if any(c.value not in (None, "") for c in reihe):
                letzte = reihe[0].row
        von = letzte + 1
        for i, reihe in enumerate(a.get("zeilen") or []):
            for si, v in enumerate(reihe, start=1):
                ws.cell(row=von + i, column=si, value=v)
        angehaengt.append({"blatt": ws.title, "von": von,
                           "bis": von + len(a.get("zeilen") or []) - 1})
    wb.calculation.fullCalcOnLoad = True
    wb.save(auftrag["ziel"])
    return {"angehaengt": angehaengt, "blaetter": [ws.title for ws in wb.worksheets]}


def _kind_arbeiten(auftrag: dict) -> dict:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return {"fehler": "bibliothek"}
    try:
        aufgabe = auftrag.get("aufgabe")
        if aufgabe == "lesen":
            return _kind_lesen(auftrag)
        if aufgabe == "anlegen":
            return _kind_anlegen(auftrag)
        if aufgabe == "aendern":
            return _kind_aendern(auftrag)
    except Fehler as e:
        return {"fehler": e.art, "meldung": str(e)}
    return {"fehler": "kaputt"}


def _kind():
    import logging
    import warnings
    logging.disable(logging.CRITICAL)
    warnings.simplefilter("ignore")
    try:
        auftrag = json.loads(sys.stdin.read() or "{}")
        antwort = _kind_arbeiten(auftrag)
    except MemoryError:
        antwort = {"fehler": "zeit"}
    except Exception:
        antwort = {"fehler": "kaputt"}
    sys.stdout.write(json.dumps(antwort, ensure_ascii=False))


if __name__ == "__main__":
    _kind()
